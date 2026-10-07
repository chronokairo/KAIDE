"""Import a company vault into ops_documents and private, checksummed Storage.

No credentials are copied to records. Source files are never modified or deleted.
"""
import argparse
from collections import Counter
from concurrent.futures import ThreadPoolExecutor, as_completed
import datetime as dt
import hashlib
from html.parser import HTMLParser
import io
import json
import logging
import mimetypes
import os
from pathlib import Path
import re
import sys
import tempfile
import urllib.error
import urllib.parse
import urllib.request

BLOCKED = {'01_sensiveis','02_contratos_assinados','03_credenciais','.git','.obsidian','.claude','node_modules','vendor','.next','dist','build','target','__pycache__'}
DENY = re.compile(r'(?i)(senha|password|credential|credencial|client[_-]?secret|recovery[_-]?code|private[_-]?key|\.env(?:\.|$)|\.pem$|\.key$|backup|\.sql$)')
SUPPORTED = {'.md','.txt','.csv','.html','.pdf','.docx','.xlsx','.xls','.png','.jpeg','.jpg','.webp','.svg','.mp4','.gsheet','.gform','.json'}
TEXT = {'.md','.txt','.csv','.html','.svg','.gsheet','.gform','.json'}
PART_BYTES = 5 * 1024 * 1024

def redact(text):
    text = text.replace('\x00', '[NUL]')
    original = text
    text = re.sub(r'-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----.*?-----END (?:RSA |EC |OPENSSH )?PRIVATE KEY-----', '[REDACTED PRIVATE KEY]', text, flags=re.S)
    text = re.sub(r'\b(?:gh[pousr]_|github_pat_|glpat-|nvapi-|sk-|sb_secret_)[A-Za-z0-9_-]{8,}', '[REDACTED]', text)
    text = re.sub(r'\beyJ[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\.[A-Za-z0-9_-]+\b', '[REDACTED JWT]', text)
    text = re.sub(r'(?i)(bearer\s+)[A-Za-z0-9._=-]+', r'\1[REDACTED]', text)
    text = re.sub(r'(?im)((?:[A-Z0-9_]*(?:API_KEY|SERVICE_ROLE_KEY|ACCESS_TOKEN)|api[_ -]?key|access[_ -]?token|password|senha|secret|segredo|seed phrase|mnemonic)\s*["\']?\s*[:=|]\s*)[^\r\n]+', r'\1[REDACTED]', text)
    text = re.sub(r'(?i)(://[^\s/:]+:)[^\s@/]+@', r'\1[REDACTED]@', text)
    return text, text != original

class PlainHTML(HTMLParser):
    def __init__(self): super().__init__(); self.parts=[]; self.ignore=0
    def handle_starttag(self, tag, attrs):
        if tag in ('script','style'): self.ignore+=1
    def handle_endtag(self, tag):
        if tag in ('script','style') and self.ignore: self.ignore-=1
    def handle_data(self, data):
        if not self.ignore: self.parts.append(data)

def extract(path):
    ext=path.suffix.lower()
    if path.stat().st_size==0:return '', 'empty_source_file'
    if ext in TEXT:
        with path.open('rb') as f: sample=f.read(2*1024*1024)
        try: text=sample.decode('utf-8-sig')
        except UnicodeDecodeError: text=sample.decode('cp1252',errors='replace')
        if ext=='.html':
            parser=PlainHTML();parser.feed(text);text='\n'.join(parser.parts)
        if ext in {'.gsheet','.gform'}:return text, 'google_shortcut_reference_only_not_exported'
        return text, 'complete' if path.stat().st_size<=len(sample) else 'preview_only_full_file_in_storage'
    if ext=='.pdf':
        from pypdf import PdfReader
        logging.getLogger('pypdf').setLevel(logging.ERROR)
        reader=PdfReader(str(path))
        if reader.is_encrypted and not reader.decrypt(''): return '', 'encrypted_pdf_original_only'
        text='\n\n'.join((page.extract_text() or '') for page in reader.pages)
        return text, 'complete' if text.strip() else 'no_text_layer_original_only'
    if ext=='.docx':
        from docx import Document
        doc=Document(str(path))
        text='\n'.join([p.text for p in doc.paragraphs]+['\t'.join(c.text for c in row.cells) for t in doc.tables for row in t.rows])
        return text, 'complete'
    if ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        out=[]
        try:
            for sheet in book:
                out.append('## '+sheet.title)
                for row in sheet.iter_rows(values_only=True):
                    if any(v is not None for v in row):out.append('\t'.join('' if v is None else str(v) for v in row))
        finally:book.close()
        return '\n'.join(out), 'complete_formulas_preserved_not_recalculated'
    if ext=='.xls':
        with path.open('rb') as source: head=source.read(2048)
        decoded=head.decode('utf-8-sig',errors='replace').lower()
        if '<table' in decoded or '<html' in decoded:
            raw=path.read_bytes()
            try:html=raw.decode('utf-8-sig')
            except UnicodeDecodeError:html=raw.decode('cp1252',errors='replace')
            parser=PlainHTML();parser.feed(html)
            return '\n'.join(parser.parts), 'complete_html_table_disguised_as_xls'
        import xlrd
        book=xlrd.open_workbook(str(path),on_demand=True)
        try:
            return '\n'.join('## '+s.name+'\n'+'\n'.join('\t'.join(str(v) for v in s.row_values(i)) for i in range(s.nrows)) for s in book.sheets()), 'complete_cached_values'
        finally:book.release_resources()
    return '', 'binary_original_only_no_ocr_or_transcription'

def inventory(root):
    files=[];skipped=[]
    for parent,dirs,names in os.walk(root):
        p=Path(parent)
        if p!=root and '.git' in dirs+names:
            skipped.append({'path':p.relative_to(root).as_posix(),'reason':'source_repository'});dirs[:]=[];continue
        for d in dirs[:]:
            if d in BLOCKED or d.startswith('.') or DENY.search(d):
                skipped.append({'path':(p/d).relative_to(root).as_posix(),'reason':'excluded_directory'});dirs.remove(d)
        for name in names:
            path=p/name;relative=path.relative_to(root).as_posix()
            if name.startswith('.') or name=='desktop.ini' or DENY.search(name):
                skipped.append({'path':relative,'reason':'excluded_name'});continue
            if path.suffix.lower() not in SUPPORTED:
                skipped.append({'path':relative,'reason':'unsupported_archive_code_or_log'});continue
            files.append(path)
    return files,skipped

def configuration(env_file, expected_ref):
    env={}
    for line in env_file.read_text().splitlines():
        match=re.match(r'^\s*(?:export\s+)?([A-Z0-9_]+)\s*=\s*(.*?)\s*$',line)
        if match:env[match[1]]=match[2].strip('"\'')
    url='https://'+expected_ref+'.supabase.co'
    if not any(v.rstrip('/')==url for k,v in env.items() if 'SUPABASE' in k and 'URL' in k):
        raise RuntimeError('The credential source does not match the expected Supabase project')
    key=env.get('SUPABASE_SERVICE_ROLE_KEY')
    if not key:raise RuntimeError('Existing service role credential not found')
    return url,key

class Backend:
    def __init__(self,url,key,bucket):self.url=url;self.key=key;self.bucket=bucket
    def call(self,method,path,body=None,headers=None,raw=False):
        data=body if isinstance(body,bytes) else None if body is None else json.dumps(body,ensure_ascii=False).encode('utf-8')
        req=urllib.request.Request(self.url+path,method=method,data=data,headers={'apikey':self.key,'Authorization':'Bearer '+self.key,'Content-Type':'application/json',**(headers or {})})
        try:
            with urllib.request.urlopen(req,timeout=45) as response:
                payload=response.read()
                return payload if raw else json.loads(payload) if payload else None
        except urllib.error.HTTPError as error:
            # Do not log response bodies, SQL values, credential-bearing requests or document contents.
            raise RuntimeError('Supabase HTTP '+str(error.code)+' on '+path.split('?')[0]) from None
    def prepare(self):
        orgs=self.call('GET','/rest/v1/organizations?select=id&slug=eq.chronokairo')
        if len(orgs)!=1:raise RuntimeError('Expected exactly one Chronokairo organization')
        self.org=orgs[0]['id']
        self.call('GET','/rest/v1/ops_documents?select=id&limit=0')
        buckets=self.call('GET','/storage/v1/bucket')
        matching=[b for b in buckets if b['id']==self.bucket]
        if matching and matching[0].get('public'):raise RuntimeError('Knowledge bucket must be private')
        if not matching:self.call('POST','/storage/v1/bucket',{'id':self.bucket,'name':self.bucket,'public':False,'file_size_limit':PART_BYTES})
    def put_parts(self,path,relative,redacted_text=None,redacted_stream=None):
        parts=[];source_hash=hashlib.sha256();stored_hash=hashlib.sha256()
        with path.open('rb') as source:
            while data:=source.read(PART_BYTES):source_hash.update(data)
        source_sha=source_hash.hexdigest()
        source_id=hashlib.sha256(relative.encode()).hexdigest()
        folder=f'{self.org}/drive/{source_id}/{source_sha}'
        stream=redacted_stream if redacted_stream is not None else io.BytesIO(redacted_text.encode('utf-8')) if redacted_text is not None else path.open('rb')
        try:
            index=0
            while data:=stream.read(PART_BYTES):
                stored_hash.update(data)
                key=f'{folder}/part-{index:05d}'
                self.call('POST','/storage/v1/object/'+self.bucket+'/'+key,data,{'Content-Type':'application/octet-stream','x-upsert':'true'})
                parts.append({'key':key,'bytes':len(data),'sha256':hashlib.sha256(data).hexdigest(),'md5':hashlib.md5(data).hexdigest()})
                index+=1
        finally:stream.close()
        listing=self.call('POST','/storage/v1/object/list/'+self.bucket,{'prefix':folder,'limit':1000,'offset':0})
        objects={o['name']:o for o in listing}
        for part in parts:
            metadata=objects.get(part['key'].split('/')[-1],{}).get('metadata') or {}
            if metadata.get('size')!=part['bytes']:raise RuntimeError('Storage size verification failed')
            etag=metadata.get('eTag') or metadata.get('etag')
            if etag and etag.strip('"')!=part['md5']:raise RuntimeError('Storage checksum verification failed')
        return {'bucket':self.bucket,'parts':parts,'source_sha256':source_sha,'stored_sha256':stored_hash.hexdigest(),'redacted_copy_only':redacted_text is not None or redacted_stream is not None}

def sync_file(backend,root,path):
    relative=path.relative_to(root).as_posix()
    try:text,extraction=extract(path)
    except Exception as error:text,extraction='', 'extraction_failed_'+type(error).__name__+'_original_preserved'
    text,changed=redact(text)
    # Full text files are redacted before upload, not just their indexed previews.
    stored_text=None
    safe_stream=None
    if path.suffix.lower() in TEXT and path.stat().st_size<=4*1024*1024:
        raw=path.read_bytes()
        try:full=raw.decode('utf-8-sig')
        except UnicodeDecodeError:full=raw.decode('cp1252',errors='replace')
        safe,full_changed=redact(full)
        if full_changed:stored_text=safe;changed=True
    elif path.suffix.lower() in TEXT:
        # Scan beyond the indexed preview. Keep the full dataset if redaction is needed.
        safe_stream=tempfile.SpooledTemporaryFile(max_size=PART_BYTES)
        any_redacted=False
        with path.open('rb') as original:
            for line in original:
                safe,line_changed=redact(line.decode('utf-8',errors='replace'))
                any_redacted|=line_changed
                safe_stream.write(safe.encode('utf-8'))
        if any_redacted:safe_stream.seek(0);changed=True
        else:safe_stream.close();safe_stream=None
    elif changed:stored_text=text
    storage=backend.put_parts(path,relative,stored_text,safe_stream)
    metadata={'source':str(root),'path':relative,'source_bytes':path.stat().st_size,'extension':path.suffix.lower(),'extraction':extraction,'secrets_redacted':changed,'storage':storage}
    body='## Origem e recuperação\n\n```json\n'+json.dumps(metadata,ensure_ascii=False,indent=2)+'\n```\n\n## Conteúdo extraído\n\n'+text
    record={'id':f'drive/{backend.org}/'+hashlib.sha256(relative.encode()).hexdigest(),'organization_id':backend.org,'title':path.stem,'category':'outro','content':body,'content_sha256':hashlib.sha256(body.encode()).hexdigest(),'document_date':None}
    backend.call('POST','/rest/v1/ops_documents?on_conflict=id',[record],{'Prefer':'resolution=merge-duplicates'})
    verified=backend.call('GET','/rest/v1/ops_documents?select=id,content_sha256&id=eq.'+urllib.parse.quote(record['id'],safe=''))
    if len(verified)!=1 or verified[0]['content_sha256']!=record['content_sha256']:raise RuntimeError('Database readback checksum failed')
    return {'path':relative,'id':record['id'],'source_bytes':metadata['source_bytes'],'parts':len(storage['parts']),'stored_bytes':sum(p['bytes'] for p in storage['parts']),'extraction':extraction,'secrets_redacted':changed,'verified':True}

def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('--root',type=Path,required=True)
    parser.add_argument('--env-file',type=Path,required=True)
    parser.add_argument('--project-ref',required=True)
    parser.add_argument('--bucket',default='chronokairo-knowledge')
    parser.add_argument('--output',type=Path,required=True)
    parser.add_argument('--apply',action='store_true')
    parser.add_argument('--retry-failed',action='store_true')
    parser.add_argument('--retry-extraction',action='store_true')
    parser.add_argument('--extra-libraries',type=Path)
    args=parser.parse_args()
    if args.extra_libraries:sys.path.insert(0,str(args.extra_libraries.resolve()))
    root=args.root.resolve()
    files,skipped=inventory(root)
    result={'started_at':dt.datetime.now(dt.timezone.utc).isoformat(),'source':str(root),'project_ref':args.project_ref,'bucket':args.bucket,'eligible_files':len(files),'skipped':skipped,'uploaded':[],'failed':[]}
    if args.retry_failed or args.retry_extraction:
        previous=json.loads(args.output.read_text(encoding='utf-8'))
        if previous['source']!=str(root) or previous['project_ref']!=args.project_ref or previous['bucket']!=args.bucket:
            raise RuntimeError('Retry manifest belongs to a different source or backend')
        wanted={r['path'] for r in previous['failed']} if args.retry_failed else {r['path'] for r in previous['uploaded'] if r['extraction'].startswith('extraction_failed')}
        files=[p for p in files if p.relative_to(root).as_posix() in wanted]
        result=previous
        if args.retry_failed:result['failed']=[]
        else:result['uploaded']=[r for r in result['uploaded'] if r['path'] not in wanted]
        result['retry_started_at']=dt.datetime.now(dt.timezone.utc).isoformat()
    args.output.parent.mkdir(parents=True,exist_ok=True)
    def save():args.output.write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    save()
    print(json.dumps({'eligible_files':len(files),'source_bytes':sum(f.stat().st_size for f in files),'skipped':len(skipped),'apply':args.apply}),flush=True)
    if not args.apply:return
    url,key=configuration(args.env_file,args.project_ref)
    backend=Backend(url,key,args.bucket);backend.prepare()
    result['organization_id']=backend.org
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures={pool.submit(sync_file,backend,root,path):path for path in files}
        for count,future in enumerate(as_completed(futures),1):
            path=futures[future]
            try:result['uploaded'].append(future.result())
            except Exception as error:result['failed'].append({'path':path.relative_to(root).as_posix(),'error':str(error) if isinstance(error,RuntimeError) else type(error).__name__})
            save()
            if count%10==0 or count==len(files):print(json.dumps({'processed':count,'uploaded':len(result['uploaded']),'failed':len(result['failed'])}),flush=True)
    result['completed_at']=dt.datetime.now(dt.timezone.utc).isoformat();save()
    print(json.dumps({'uploaded':len(result['uploaded']),'failed':len(result['failed']),'stored_bytes':sum(r['stored_bytes'] for r in result['uploaded']),'extraction_statuses':Counter(r['extraction'] for r in result['uploaded'])}),flush=True)
    if result['failed']:sys.exit(1)

if __name__=='__main__':main()
