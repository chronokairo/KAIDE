"""Separate workbook prospects into person (B2C) and business (B2B) CRM records."""
import argparse, hashlib, json, re, sys, unicodedata
from pathlib import Path
from materialize_knowledge import ResilientBackend, configuration, identifier, metadata

def person_key(name):
    return re.sub(r'[^a-z0-9]','',unicodedata.normalize('NFKD',name.casefold()).encode('ascii','ignore').decode())
def valid_person(name):
    if not isinstance(name,str):return False
    return len(name.split())>=2 and not re.search(r'\d|\b(TESTE|ALUNO|TOTAL|TURMA|PER[IÍ]ODO|CURSO)\b',name,re.I)
def candidates(report,docs,recovered):
    people={}
    def add(name,email,source,sheet,row):
        if not valid_person(name):return
        key=person_key(name);person=people.setdefault(key,{'name':name.strip(),'email':None,'sources':[]})
        if isinstance(email,str) and re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',email.strip()):person['email']=email.strip().lower()
        person['sources'].append({'source_document_id':source['id'],'path':source['path'],'sheet':sheet,'row':row})
    for e in report['uploaded']:
        if '/leads/cursos_ia_python/' not in e['path'] or not e['path'].lower().endswith(('.xlsx','.xls')):continue
        meta=metadata(docs[e['id']]['content']);p=recovered/(e['id'].split('/')[-1]+meta['extension'])
        if hashlib.sha256(p.read_bytes()).hexdigest()!=meta['storage']['stored_sha256']:raise ValueError('Workbook source checksum mismatch')
        if p.suffix.lower()=='.xlsx':
            from openpyxl import load_workbook
            book=load_workbook(p,read_only=True,data_only=True)
            try:
                for sheet in book:
                    started=False
                    for index,row in enumerate(sheet.iter_rows(values_only=True),1):
                        value=row[0] if row else None
                        if value=='Nome':started=True;continue
                        if started:add(value,None,e,sheet.title,index)
            finally:book.close()
        else:
            import xlrd
            book=xlrd.open_workbook(str(p),on_demand=True)
            try:
                for sheet in book.sheets():
                    for index in range(sheet.nrows):
                        row=sheet.row_values(index)
                        if len(row)>4 and row[3]:add(row[3],row[4],e,sheet.name,index+1)
            finally:book.release_resources()
    return people

def classify(existing):
    try:notes=json.loads(existing.get('notes') or '{}')
    except ValueError:notes={}
    if isinstance(notes,dict) and notes.get('market') in {'B2B','B2C'}:return notes['market']
    digits=re.sub(r'\D','',existing.get('tax_id') or '')
    if len(digits)==14 or existing.get('company'):return 'B2B'
    if len(digits)==11:return 'B2C'
    return 'unclassified'

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--report',type=Path,required=True);parser.add_argument('--documents',type=Path,required=True);parser.add_argument('--recovered',type=Path,required=True);parser.add_argument('--env-file',type=Path,required=True);parser.add_argument('--extra-libraries',type=Path);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
    if args.extra_libraries:sys.path.insert(0,str(args.extra_libraries.resolve()))
    report=json.loads(args.report.read_text(encoding='utf8'));docs={d['id']:d for d in json.loads(args.documents.read_text(encoding='utf8'))};people=candidates(report,docs,args.recovered)
    url,key=configuration(args.env_file,report['project_ref']);b=ResilientBackend(url,key,report['bucket']);b.org=report['organization_id']
    existing=b.call('GET','/rest/v1/leads?organization_id=eq.'+b.org+'&select=*&limit=1000');by_name={person_key(x['name']):x for x in existing if classify(x)!='B2B'};created=[];linked=[]
    for key,p in people.items():
        original=by_name.get(key)
        try:notes=json.loads(original.get('notes') or '{}') if original else {}
        except ValueError:notes={'previous_notes':original['notes']}
        if not isinstance(notes,dict):notes={'previous_notes':notes}
        notes.update(person_type='PF',market='B2C',source_kind='course_roster',sources=p['sources'])
        if original:
            changes={'notes':json.dumps(notes,ensure_ascii=False)}
            if not original.get('email') and p['email']:changes['email']=p['email']
            b.call('PATCH','/rest/v1/leads?id=eq.'+original['id']+'&organization_id=eq.'+b.org,changes);linked.append(original['id'])
        else:
            rid=identifier(b.org,'person-lead',key)
            b.call('POST','/rest/v1/leads',[{'id':rid,'organization_id':b.org,'name':p['name'],'company':None,'email':p['email'],'segment':'Pessoa física — cursos de IA/Python','source':'Evento','status':'Novo','notes':json.dumps(notes,ensure_ascii=False)}]);created.append(rid)
    for x in existing:
        if classify(x)!='B2B':continue
        try:notes=json.loads(x.get('notes') or '{}')
        except ValueError:notes={'previous_notes':x['notes']}
        if not isinstance(notes,dict):notes={'previous_notes':notes}
        notes.update(person_type='PJ',market='B2B')
        b.call('PATCH','/rest/v1/leads?id=eq.'+x['id']+'&organization_id=eq.'+b.org,{'notes':json.dumps(notes,ensure_ascii=False)})
    saved=b.call('GET','/rest/v1/leads?organization_id=eq.'+b.org+'&select=*&limit=1000');groups={k:[x for x in saved if classify(x)==k] for k in ['B2B','B2C','unclassified']}
    if not set(created+linked).issubset({x['id'] for x in groups['B2C']}):raise ValueError('Person classification readback failed')
    args.output.mkdir(parents=True,exist_ok=True)
    for market,rows in groups.items():(args.output/('leads-'+market+'.json')).write_text(json.dumps(rows,ensure_ascii=False,indent=2),encoding='utf8')
    result={'new_person_leads':len(created),'existing_person_leads_linked':len(linked),'unique_people_in_workbooks':len(people),'counts':{k:len(v) for k,v in groups.items()},'source_rows':sum(len(p['sources']) for p in people.values()),'people_with_email':sum(bool(p['email']) for p in people.values())}
    (args.output/'person-lead-report.json').write_text(json.dumps(result,indent=2),encoding='utf8');print(json.dumps(result))
if __name__=='__main__':main()
