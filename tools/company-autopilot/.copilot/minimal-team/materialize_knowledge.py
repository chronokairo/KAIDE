"""Recover imported Storage objects and materialize provenance-linked application data.

Run against an import report; never deletes source files. Cleanup uses the separate
PowerShell verifier and requires an explicit --apply. Credentials stay in memory.
"""
import argparse, csv, datetime, hashlib, io, json, re, sys, tempfile, uuid, time, urllib.error
from pathlib import Path
from concurrent.futures import ThreadPoolExecutor, as_completed
from html.parser import HTMLParser
from urllib.parse import quote
from sync_knowledge import Backend, configuration, redact

NAMESPACE=uuid.UUID('a60e5577-a521-4897-9629-a4d843a48afc')
class ResilientBackend(Backend):
    def call(self,*args,**kwargs):
        for attempt in range(5):
            try:return super().call(*args,**kwargs)
            except (urllib.error.URLError,TimeoutError,ConnectionError):
                if attempt==4:raise
                time.sleep(min(attempt+1,4))
def identifier(org,kind,key):return str(uuid.uuid5(NAMESPACE,org+'/'+kind+'/'+key))
def normalized(value):return re.sub(r'\W','',value.casefold())
def safe_value(value):
    if isinstance(value,str):return redact(value)[0]
    if isinstance(value,dict):return {k:safe_value(v) for k,v in value.items()}
    if isinstance(value,(list,tuple)):return [safe_value(v) for v in value]
    return value
def metadata(content):return json.loads(content.split('```json\n',1)[1].split('\n```',1)[0])
def area(path):
    lower=path.lower()
    for key,name in [('financeiro','finance'),('comercial','crm'),('leads','crm'),('07_clientes','projects'),('06_socios','people'),('06_funcionarios','people'),('marketing','marketing'),('engenharia','engineering')]:
        if key in lower:return name
    return 'operations'
def markdown_tables(text):
    result=[];headers=None
    for line in text.splitlines():
        if not line.strip().startswith('|'):headers=None;continue
        line=re.sub(r'\[\[.*?\]\]',lambda m:m.group().replace('|','\u0001'),line)
        cells=[c.strip().replace('\u0001','|') for c in line.strip().strip('|').split('|')]
        if all(re.fullmatch(r'[: -]+',c or ' ') for c in cells):continue
        if headers is None:headers=cells;continue
        if len(cells)==len(headers):result.append(dict(zip(headers,cells)))
    return result
class TableHTML(HTMLParser):
    def __init__(self):super().__init__();self.rows=[];self.row=[];self.cell=None
    def handle_starttag(self,tag,attrs):
        if tag=='tr':self.row=[]
        if tag in ('td','th'):self.cell=[]
    def handle_data(self,data):
        if self.cell is not None:self.cell.append(data)
    def handle_endtag(self,tag):
        if tag in ('td','th') and self.cell is not None:self.row.append(' '.join(self.cell).strip());self.cell=None
        if tag=='tr' and self.row:self.rows.append(self.row)

def records(path,extracted):
    ext=path.suffix.lower()
    if ext=='.csv':
        with path.open('r',encoding='utf-8-sig',errors='replace',newline='') as f:
            sample=f.read(10000);f.seek(0)
            try:dialect=csv.Sniffer().sniff(sample,delimiters=',;\t|')
            except csv.Error:dialect=csv.excel
            yield from ({'kind':'row','values':r} for r in csv.DictReader(f,dialect=dialect))
    elif ext=='.xlsx':
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False)
        try:
            for sheet in book:
                for index,row in enumerate(sheet.iter_rows(values_only=True),1):
                    if any(v is not None for v in row):yield {'kind':'spreadsheet_row','sheet':sheet.title,'row_number':index,'values':list(row),'formulas_preserved':True}
        finally:book.close()
    elif ext=='.xls':
        raw=path.read_bytes()
        if b'<table' in raw[:2048].lower() or b'<html' in raw[:2048].lower():
            p=TableHTML();p.feed(raw.decode('utf-8-sig',errors='replace'))
            for i,row in enumerate(p.rows,1):yield {'kind':'spreadsheet_row','row_number':i,'values':row}
        else:
            import xlrd
            book=xlrd.open_workbook(str(path),on_demand=True)
            try:
                for sheet in book.sheets():
                    for i in range(sheet.nrows):yield {'kind':'spreadsheet_row','sheet':sheet.name,'row_number':i+1,'values':sheet.row_values(i)}
            finally:book.release_resources()
    elif extracted.strip():
        for r in markdown_tables(extracted):yield {'kind':'table_row','values':r}
        for i,section in enumerate(re.split(r'(?m)(?=^#{1,6} )',extracted)):
            if section.strip():yield {'kind':'knowledge_section','section_number':i,'text':section.strip()}
    else:yield {'kind':'asset','media_type':ext,'requires_ocr_or_transcription':ext in {'.pdf','.png','.jpg','.jpeg','.mp4'}}

def recover(b,entry,doc,output):
    meta=metadata(doc['content']);s=meta['storage'];p=output/'recovered'/entry['id'].split('/')[-1]
    p=p.with_suffix(meta['extension']);p.parent.mkdir(parents=True,exist_ok=True)
    if p.exists() and hashlib.sha256(p.read_bytes()).hexdigest()==s['stored_sha256']:return entry,doc,meta,p
    h=hashlib.sha256()
    with p.open('wb') as f:
        for part in s['parts']:
            raw=b.call('GET','/storage/v1/object/'+s['bucket']+'/'+quote(part['key'],safe='/'),raw=True)
            if len(raw)!=part['bytes'] or hashlib.sha256(raw).hexdigest()!=part['sha256']:raise RuntimeError('Recovered part checksum mismatch')
            f.write(raw);h.update(raw)
    if h.hexdigest()!=s['stored_sha256']:raise RuntimeError('Recovered source checksum mismatch')
    return entry,doc,meta,p

def materialize(b,item,output):
    entry,doc,meta,p=item;target=area(entry['path']);source_id=entry['id'].split('/')[-1]
    extracted=doc['content'].split('## Conteúdo extraído\n\n',1)[-1]
    parts=[];count=0;size=0;buffer=bytearray();columns=set()
    def flush():
        nonlocal buffer,size
        if not buffer:return
        raw=bytes(buffer);key=f'{b.org}/application-data/{source_id}/{meta["storage"]["stored_sha256"]}/part-{len(parts):05d}-{hashlib.sha256(raw).hexdigest()}.ndjson'
        b.call('POST','/storage/v1/object/'+b.bucket+'/'+key,raw,{'Content-Type':'application/x-ndjson','x-upsert':'true'})
        check=b.call('GET','/storage/v1/object/'+b.bucket+'/'+key,raw=True)
        if check!=raw:raise RuntimeError('Structured dataset readback mismatch')
        parts.append({'key':key,'bytes':len(raw),'sha256':hashlib.sha256(raw).hexdigest()});size+=len(raw);buffer=bytearray()
    for record in records(p,extracted):
        if isinstance(record.get('values'),dict):columns.update(k for k in record['values'] if k)
        record.update(source_document_id=entry['id'],source_path=entry['path'],application=target)
        line=json.dumps(safe_value(record),ensure_ascii=False,default=str)+'\n';raw=line.encode('utf8')
        if len(raw)>5*1024*1024:raise RuntimeError('Structured row exceeds bucket limit')
        if len(buffer)+len(raw)>5*1024*1024:flush()
        buffer.extend(raw);count+=1
    flush()
    descriptor={'version':1,'organization_id':b.org,'source_document_id':entry['id'],'source_path':entry['path'],'application':target,'format':'ndjson','record_count':count,'bucket':b.bucket,'parts':parts,'source_sha256':meta['storage']['source_sha256'],'facts_require_review':True}
    inv={'id':identifier(b.org,'inventory',entry['id']),'organization_id':b.org,'process_name':target+': '+entry['path'],'category':'Importação de fonte','data_types':sorted(columns) or ['knowledge' if extracted.strip() else 'media'],'purpose':json.dumps(descriptor,ensure_ascii=False),'security_measures':'Private Storage; organization RLS; source provenance; no automatic financial posting'}
    b.call('POST','/rest/v1/data_inventory?on_conflict=id',[inv],{'Prefer':'resolution=merge-duplicates'})
    saved=b.call('GET','/rest/v1/data_inventory?id=eq.'+inv['id']+'&select=id,purpose')
    if len(saved)!=1 or saved[0]['purpose']!=inv['purpose']:raise RuntimeError('Inventory readback mismatch')
    category='comercial' if target=='crm' else 'relatorio' if any(w in entry['path'].lower() for w in ['relatorio','diagnostico']) else 'registro'
    b.call('PATCH','/rest/v1/ops_documents?id=eq.'+quote(entry['id'],safe=''),{'category':category})
    return descriptor

def native_crm(b,items):
    existing=b.call('GET','/rest/v1/leads?organization_id=eq.'+b.org+'&select=*&limit=1000');taxes={normalized(x.get('tax_id') or '') for x in existing};names={normalized(x['name']) for x in existing};created=[]
    for entry,doc,meta,p in items:
        if entry['path']!='04_dados_externos/leads_chronokairo.xlsx':continue
        from openpyxl import load_workbook
        book=load_workbook(p,read_only=True,data_only=True)
        try:
            sheet=book['Pipeline'];rows=sheet.iter_rows(values_only=True);headers=next(rows)
            for values in rows:
                r=dict(zip(headers,values));name=r.get('Empresa');tax=str(r.get('CNPJ') or '')
                if not name or normalized(tax) in taxes or normalized(name) in names:continue
                clean=lambda v:None if v in (None,'','-') else str(v)
                record={'id':identifier(b.org,'lead',tax or name),'organization_id':b.org,'name':name,'company':name,'tax_id':clean(tax),'email':clean(r.get('E-mail')),'phone':clean(r.get('Telefone')),'segment':clean(r.get('Segmento')),'tier':r.get('Tier') if r.get('Tier') in {'A','B','C'} else None,'location':clean(r.get('Cidade/UF')),'source':'Outbound','status':'Novo','approach':clean(r.get('Proxima acao')),'notes':json.dumps({'source_document_id':entry['id'],'original_status':r.get('Status'),'source_row':r},ensure_ascii=False,default=str)}
                b.call('POST','/rest/v1/leads',[record]);saved=b.call('GET','/rest/v1/leads?id=eq.'+record['id']+'&select=id,tax_id')
                if len(saved)!=1 or saved[0]['tax_id']!=record['tax_id']:raise RuntimeError('CRM readback mismatch')
                created.append(record['id']);names.add(normalized(name));taxes.add(normalized(tax))
        finally:book.close()
    return created

def native_clients(b,items):
    customers=b.call('GET','/rest/v1/customers?organization_id=eq.'+b.org+'&select=*&limit=1000')
    names={normalized(x['name']):x['id'] for x in customers};projects=b.call('GET','/rest/v1/projects?organization_id=eq.'+b.org+'&select=*&limit=1000');codes={x.get('code') for x in projects};result={'customers':[],'projects':[]}
    for entry,doc,meta,p in items:
        if entry['path']!='07_clientes/README.md':continue
        text=doc['content'].split('## Conteúdo extraído\n\n',1)[-1].split('## 🔴',1)[0]
        for row in markdown_tables(text):
            name=row.get('Cliente','').replace('**','').strip()
            # Only explicitly named current accounts absent from the existing CRM.
            if name not in {'Dailson Queiroz','Evolutvs'}:continue
            key=normalized(name);cid=names.get(key)
            provenance=json.dumps({'source_document_id':entry['id'],'source_row':row,'as_of':'2026-09-19'},ensure_ascii=False)
            if not cid:
                cid=identifier(b.org,'customer',key)
                b.call('POST','/rest/v1/customers',[{'id':cid,'organization_id':b.org,'name':name,'segmento':row.get('Segmento'),'observacoes':provenance,'origem_cadastro':'Drive knowledge import'}]);result['customers'].append(cid);names[key]=cid
            code='DRIVE-'+key.upper()
            if code not in codes:
                pid=identifier(b.org,'project',code)
                b.call('POST','/rest/v1/projects',[{'id':pid,'organization_id':b.org,'name':name,'customer_id':cid,'code':code,'description':row.get('Resumo do Escopo','')+'\n'+provenance}]);result['projects'].append(pid);codes.add(code)
    for table in ('customers','projects'):
        for rid in result[table]:
            if len(b.call('GET','/rest/v1/'+table+'?id=eq.'+rid+'&select=id'))!=1:raise RuntimeError('Client/project readback mismatch')
    return result

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--report',type=Path,required=True);parser.add_argument('--env-file',type=Path,required=True);parser.add_argument('--output',type=Path,required=True);parser.add_argument('--extra-libraries',type=Path);args=parser.parse_args()
    if args.extra_libraries:sys.path.insert(0,str(args.extra_libraries.resolve()))
    report=json.loads(args.report.read_text(encoding='utf8'));url,key=configuration(args.env_file,report['project_ref']);b=ResilientBackend(url,key,report['bucket']);b.prepare()
    docs={x['id']:x for x in b.call('GET','/rest/v1/ops_documents?organization_id=eq.'+b.org+'&select=id,title,content&limit=1000')};args.output.mkdir(parents=True,exist_ok=True)
    result={'started_at':datetime.datetime.now(datetime.timezone.utc).isoformat(),'organization_id':b.org,'datasets':[],'failed':[],'cleanup_eligible':[]};items=[]
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures={pool.submit(recover,b,e,docs[e['id']],args.output):e for e in report['uploaded']}
        for f in as_completed(futures):
            e=futures[f]
            try:items.append(f.result())
            except Exception as error:result['failed'].append({'path':e['path'],'error':str(error)})
    print(json.dumps({'recovered':len(items),'failed':len(result['failed'])}),flush=True)
    result['created_leads']=native_crm(b,items)
    result['created_client_records']=native_clients(b,items)
    with ThreadPoolExecutor(max_workers=4) as pool:
        futures={pool.submit(materialize,b,item,args.output):item for item in items}
        for f in as_completed(futures):
            item=futures[f];e,doc,meta,p=item
            try:
                result['datasets'].append(f.result());s=meta['storage']
                if not s['redacted_copy_only'] and s['source_sha256']==s['stored_sha256']:
                    result['cleanup_eligible'].append({'path':e['path'],'sha256':s['source_sha256'],'source_bytes':e['source_bytes'],'storage_download_verified':True})
            except Exception as error:result['failed'].append({'path':e['path'],'error':str(error)})
            if len(result['datasets'])%25==0:print(json.dumps({'datasets':len(result['datasets']),'failed':len(result['failed'])}),flush=True)
            (args.output/'materialization-report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8')
    result['completed_at']=datetime.datetime.now(datetime.timezone.utc).isoformat();(args.output/'materialization-report.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf8');print(json.dumps({'datasets':len(result['datasets']),'new_leads':len(result['created_leads']),'cleanup_eligible':len(result['cleanup_eligible']),'failed':len(result['failed'])}))
if __name__=='__main__':main()
